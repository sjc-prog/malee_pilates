from pathlib import Path
from datetime import date

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path("/Users/jcohen/Work/malee-pilates")
OUT_DIR = ROOT / "docs"
OUT_DIR.mkdir(parents=True, exist_ok=True)
OUT = OUT_DIR / "Malee-Pilates-VMS-Master-Data.docx"
LOGO = ROOT / "site/img/email/logo.png"

INK = "45403A"
SAGE = "4E5A3C"
PALE_SAGE = "EEF3EA"
BLUSH = "E9B9C5"
PALE_BLUSH = "F9EDF0"
CREAM = "FFFAF7"
LIGHT_GREY = "D9D9D9"
MID_GREY = "6F6A64"
WHITE = "FFFFFF"


def set_cell_fill(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=110, bottom=100, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=LIGHT_GREY, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def keep_row_together(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_column_width(cell, width_cm):
    cell.width = Cm(width_cm)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_cm * 567)))
    tc_w.set(qn("w:type"), "dxa")


def style_run(run, size=9.2, bold=False, color=INK, font="Aptos"):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)


def add_table(doc, headers, rows, widths=None, font_size=8.6, first_col_bold=False):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    header = table.rows[0]
    set_repeat_table_header(header)
    for i, label in enumerate(headers):
        cell = header.cells[i]
        set_cell_fill(cell, SAGE)
        set_cell_margins(cell, 120, 110, 120, 110)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(str(label))
        style_run(run, font_size, True, WHITE)
        if widths:
            set_column_width(cell, widths[i])
    for r_index, values in enumerate(rows):
        row = table.add_row()
        keep_row_together(row)
        for i, value in enumerate(values):
            cell = row.cells[i]
            set_cell_fill(cell, WHITE if r_index % 2 == 0 else PALE_SAGE)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            if i > 0 and len(str(value)) < 18:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(str(value))
            style_run(run, font_size, first_col_bold and i == 0, INK)
            if widths:
                set_column_width(cell, widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_bullet(doc, text, level=0, bold_lead=None):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(2)
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        style_run(r1, 9.6, True)
        r2 = p.add_run(text[len(bold_lead):])
        style_run(r2, 9.6)
    else:
        r = p.add_run(text)
        style_run(r, 9.6)
    return p


def add_note(doc, lead, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(7)
    r1 = p.add_run(lead + " ")
    style_run(r1, 9.6, True, SAGE)
    r2 = p.add_run(text)
    style_run(r2, 9.6, False, INK)


doc = Document()
section = doc.sections[0]
section.page_width = Cm(21.0)
section.page_height = Cm(29.7)
section.top_margin = Cm(1.55)
section.bottom_margin = Cm(1.45)
section.left_margin = Cm(1.55)
section.right_margin = Cm(1.55)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal.font.size = Pt(9.8)
normal.font.color.rgb = RGBColor.from_string(INK)
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.line_spacing = 1.12

for style_name, size, before, after in (
    ("Title", 25, 0, 6),
    ("Subtitle", 11, 0, 14),
    ("Heading 1", 16, 16, 7),
    ("Heading 2", 12.5, 12, 5),
    ("Heading 3", 10.5, 8, 4),
):
    style = styles[style_name]
    style.font.name = "Aptos Display" if style_name in ("Title", "Heading 1") else "Aptos"
    style.font.size = Pt(size)
    style.font.bold = style_name != "Subtitle"
    style.font.color.rgb = RGBColor.from_string("000000")
    style._element.rPr.rFonts.set(qn("w:eastAsia"), style.font.name)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

# Footer
footer = section.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
fr = fp.add_run("MÀLEE Pilates  •  VMS Master Data  •  Internal working document")
style_run(fr, 8.0, False, MID_GREY)

# Cover
if LOGO.exists():
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(15)
    run = p.add_run()
    run.add_picture(str(LOGO), width=Inches(2.0))

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run("MÀLEE Pilates VMS Master Data")
title_ppr = title._p.get_or_add_pPr()
title_border = OxmlElement("w:pBdr")
title_bottom = OxmlElement("w:bottom")
title_bottom.set(qn("w:val"), "nil")
title_border.append(title_bottom)
title_ppr.append(title_border)

subtitle = doc.add_paragraph(style="Subtitle")
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.add_run("Services  Classes  Schedule  Trainers  Pricing  Launch Promotions")

meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
meta.paragraph_format.space_before = Pt(8)
meta.paragraph_format.space_after = Pt(20)
r = meta.add_run("Prepared 16 October 2026  •  Working source of truth for back-office setup")
style_run(r, 9.2, False, MID_GREY)

intro = doc.add_paragraph()
intro.paragraph_format.space_after = Pt(8)
r = intro.add_run(
    "This document consolidates the latest client-supplied package lists, the regular Bophut timetable, "
    "the Nathon private-service details, the Halloween soft-opening event, and the Bophut instructor policy. "
    "Items marked Ready can be entered into the VMS. Items marked Confirm should be held until the client resolves the stated point."
)
style_run(r, 10.0)

add_table(
    doc,
    ["Setup area", "Count", "Status"],
    [
        ("Regular group class types", "6", "Ready"),
        ("Weekly recurring group sessions", "32", "Ready"),
        ("Current trainers", "3", "Ready"),
        ("Group purchase options", "6", "Ready"),
        ("Current private purchase options", "3 + add-on", "Ready"),
        ("Halloween launch classes", "5", "Ready as one-day override"),
        ("Legacy private offers", "2", "Confirm before creating"),
    ],
    widths=[9.2, 3.0, 5.3],
    font_size=9.0,
    first_col_bold=True,
)

doc.add_page_break()

doc.add_heading("1 Venue and Area Structure", level=1)
add_table(
    doc,
    ["Venue or area", "Use", "Capacity", "Booking pattern", "Status"],
    [
        (
            "Bophut Reformer Studio",
            "Regular group classes and private sessions",
            "9 group places",
            "Recurring timetable; private availability 12:45-17:25 by appointment",
            "Ready; retain existing VMS address",
        ),
        (
            "Nathon Hillside Sea View and Pier Studio",
            "Private Reformer only",
            "6 reformers; 1-6 people",
            "Advance booking at 09:30 or 11:00, subject to availability",
            "Ready",
        ),
    ],
    widths=[3.9, 4.3, 2.5, 4.8, 3.0],
    font_size=8.2,
    first_col_bold=True,
)
add_note(
    doc,
    "Nathon location",
    "Soi Mdet, Lipa Yai, Nathon. The source describes this as a peaceful hillside studio with sea and pier views.",
)
add_note(
    doc,
    "Recommended VMS areas",
    "Create a Bophut Reformer Studio area with capacity 9 and a separate Nathon Private Reformer Studio area with capacity 6.",
)

doc.add_heading("2 Service and Class Catalogue", level=1)
doc.add_heading("Regular Group Classes", level=2)
add_table(
    doc,
    ["Class", "Level", "Duration", "Capacity", "Customer description"],
    [
        ("Morning Glow", "All levels", "55 min", "9", "Morning core activation and refreshing energy."),
        ("Full Body Flow", "All levels", "55 min", "9", "Total-body toning and endurance."),
        ("Beginner", "Beginner", "55 min", "9", "First-timer-friendly foundations, safe form and posture."),
        ("Booty and Core", "All levels", "55 min", "9", "Focused glute and core work."),
        ("Sunset Tone", "All levels", "55 min", "9", "De-stress, mobility and body sculpting after the working day."),
        ("Jumpboard", "Not stated", "55 min", "9", "Low-impact cardio and energetic reformer work."),
    ],
    widths=[3.2, 2.6, 2.1, 2.0, 8.3],
    font_size=8.5,
    first_col_bold=True,
)

doc.add_heading("Private Services", level=2)
add_table(
    doc,
    ["Service", "Area", "Duration", "Capacity", "Notes"],
    [
        ("Private Reformer 1 to 1", "Bophut or Nathon", "60 min provisional", "1", "Session tailored to goals, level and individual needs."),
        ("Private Reformer Group", "Nathon", "60 min provisional", "2-6", "Private session for family or friends; additional-person pricing applies."),
        ("Private Reformer Appointment", "Bophut", "Confirm", "1-6", "Available within the 12:45-17:25 timetable window by appointment."),
    ],
    widths=[4.0, 3.2, 3.0, 2.0, 6.0],
    font_size=8.5,
    first_col_bold=True,
)
add_note(doc, "Duration note", "The older private-apparatus menu specifies one-hour sessions. The newest private price message does not restate the duration, so the VMS duration should be confirmed before publication.")

doc.add_heading("3 Regular Weekly Schedule", level=1)
add_note(doc, "Schedule rule", "Every listed recurring group class lasts 55 minutes. Private sessions may be booked by appointment between 12:45 and 17:25 where availability allows.")

schedule = {
    "Monday": [
        ("09:00-09:55", "Full Body Flow", "Kru Yui"),
        ("10:15-11:10", "Booty and Core", "Kru Yui"),
        ("11:30-12:25", "Full Body Flow", "Kru Yui"),
        ("17:30-18:25", "Sunset Tone", "Kru Yui"),
    ],
    "Tuesday": [
        ("07:00-07:55", "Morning Glow", "Kru Yui"),
        ("09:00-09:55", "Booty and Core", "Kru Yui"),
        ("10:15-11:10", "Beginner", "Kru Yui"),
        ("11:30-12:25", "Full Body Flow", "Kru Yui"),
        ("17:30-18:25", "Beginner", "Kru Leaw"),
        ("18:45-19:40", "Booty and Core", "Kru Leaw"),
    ],
    "Wednesday": [
        ("09:00-09:55", "Full Body Flow", "Kru Nex"),
        ("10:15-11:10", "Jumpboard", "Kru Nex"),
        ("11:30-12:25", "Beginner", "Kru Nex"),
        ("17:30-18:25", "Sunset Tone", "Kru Nex"),
    ],
    "Thursday": [
        ("07:00-07:55", "Morning Glow", "Kru Nex"),
        ("09:00-09:55", "Booty and Core", "Kru Nex"),
        ("10:15-11:10", "Beginner", "Kru Nex"),
        ("11:30-12:25", "Full Body Flow", "Kru Nex"),
        ("17:30-18:25", "Beginner", "Kru Nex"),
        ("18:45-19:40", "Full Body Flow", "Kru Nex"),
    ],
    "Friday": [
        ("09:00-09:55", "Jumpboard", "Kru Nex"),
        ("10:15-11:10", "Full Body Flow", "Kru Nex"),
        ("11:30-12:25", "Booty and Core", "Kru Nex"),
        ("17:30-18:25", "Sunset Tone", "Kru Nex"),
    ],
    "Saturday": [
        ("09:00-09:55", "Full Body Flow", "Kru Nex"),
        ("10:15-11:10", "Jumpboard", "Kru Nex"),
        ("11:30-12:25", "Beginner", "Kru Nex"),
        ("17:30-18:25", "Sunset Tone", "Kru Nex"),
        ("18:45-19:40", "Booty and Core", "Kru Nex"),
    ],
    "Sunday": [
        ("09:00-09:55", "Beginner", "Kru Leaw"),
        ("10:15-11:10", "Full Body Flow", "Kru Leaw"),
        ("11:30-12:25", "Booty and Core", "Kru Leaw"),
    ],
}

for day, rows in schedule.items():
    doc.add_heading(day, level=2)
    add_table(
        doc,
        ["Time", "Class", "Trainer", "Duration"],
        [(time, class_name, trainer, "55 min") for time, class_name, trainer in rows],
        widths=[3.4, 7.2, 4.3, 3.0],
        font_size=8.7,
        first_col_bold=False,
    )

doc.add_heading("4 Trainers", level=1)
add_table(
    doc,
    ["Trainer", "Regular teaching allocation", "Launch event allocation", "VMS note"],
    [
        ("Kru Yui", "Monday and Tuesday daytime; 8 weekly sessions", "Ghostly Halloween Flow", "Create or retain active trainer profile"),
        ("Kru Nex", "Wednesday to Saturday; 19 weekly sessions", "Spooktacular Booty and Core; Witch's First Step; Nightmare Night", "Create or retain active trainer profile"),
        ("Kru Leaw", "Tuesday evening and Sunday morning; 5 weekly sessions", "First Bite Beginner Flow", "Create or retain active trainer profile"),
    ],
    widths=[2.7, 5.3, 6.1, 4.1],
    font_size=8.3,
    first_col_bold=True,
)
add_note(doc, "Missing trainer data", "The supplied files do not contain bios, qualifications, contact details, employment status, or profile descriptions.")

doc.add_heading("5 Group Class Pricing", level=1)
add_note(doc, "Source authority", "Use this table as the current VMS configuration. The client labelled the underlying message as final, so it overrides the older Pilates Princess promotion PDF.")
add_table(
    doc,
    ["Product", "Price", "Included", "Effective rate", "Validity", "Sharing"],
    [
        ("Introduction Offer", "THB 1,800", "3 classes", "THB 600", "14 days from purchase", "None; first-time clients only"),
        ("Drop-in Class", "THB 900", "1 class", "THB 900", "Booked day only", "None"),
        ("5-Class Package", "THB 4,000", "5 classes", "THB 800", "30 days from purchase", "Up to 1 class"),
        ("10-Class Package", "THB 7,500", "10 classes", "THB 750", "60 days from purchase", "Up to 3 classes"),
        ("20-Class Package", "THB 14,000", "20 classes", "THB 700", "90 days from purchase", "Up to 5 classes"),
        ("30-Class Package", "THB 19,500", "30 classes", "THB 650", "120 days from purchase", "Up to 15 classes"),
    ],
    widths=[3.6, 2.7, 2.5, 2.5, 3.6, 3.6],
    font_size=8.2,
    first_col_bold=True,
)
add_note(doc, "Sharing rule", "The package owner must attend the same class. Up to two friends may join a class, subject to the total sharing allowance of the package.")

doc.add_heading("6 Private Pricing", level=1)
add_table(
    doc,
    ["Product", "Price", "Effective rate", "Validity", "Eligibility and notes"],
    [
        ("Private Trial Class", "THB 1,800", "THB 1,800", "Single session", "First-time private client"),
        ("Private Drop-in", "THB 2,500", "THB 2,500", "Single session", "Flexible one-off booking"),
        ("5 Private Sessions", "THB 11,500", "THB 2,300", "30 days", "Current private package"),
        ("Additional Participant", "+THB 1,000", "Per person per session", "Same booked session", "Maximum total class size 6"),
    ],
    widths=[4.0, 3.0, 3.2, 2.7, 5.5],
    font_size=8.5,
    first_col_bold=True,
)
add_note(doc, "Current configuration", "Do not create the legacy 10 Private Sessions package or the older two-person group-private promotion until the client confirms they remain on sale.")

doc.add_heading("7 Soft Opening Package Gift", level=1)
add_bullet(doc, "Group packages: applies to the 5-, 10-, 20-, and 30-class packages.")
add_bullet(doc, "Private packages: applies to the 5 Private Sessions package.")
add_bullet(doc, "Gift: one complimentary 60-minute Thai massage or foot massage.")
add_bullet(doc, "Redemption: within 14 days of package purchase.")
add_bullet(doc, "Exclusions: group Introduction Offer, group Drop-in, Private Trial, and Private Drop-in.")
add_bullet(doc, "Availability: limited quantity; the exact stock allocation was not supplied.")

doc.add_page_break()
doc.add_heading("8 Halloween Soft Opening Event", level=1)
intro = doc.add_paragraph()
r = intro.add_run("Ghost Princess Pilates Soft Opening Halloween Free Class")
style_run(r, 11.0, True, SAGE)
intro.add_run(" takes place on Saturday 31 October 2026. All five sessions are free and require online booking.")

add_table(
    doc,
    ["Time", "Class name", "Trainer", "Duration", "Price", "Capacity"],
    [
        ("09:00-09:55", "Ghostly Halloween Flow", "Kru Yui", "55 min", "THB 0", "9"),
        ("10:15-11:10", "Spooktacular Booty and Core", "Kru Nex", "55 min", "THB 0", "9"),
        ("11:30-12:25", "First Bite Beginner Flow", "Kru Leaw", "55 min", "THB 0", "9"),
        ("17:30-18:25", "Witch's First Step Beginner", "Kru Nex", "55 min", "THB 0", "9"),
        ("18:45-19:45", "Nightmare Night Halloween Flow", "Kru Nex", "60 min", "THB 0", "9"),
    ],
    widths=[3.0, 6.0, 3.2, 2.2, 2.0, 2.2],
    font_size=8.4,
    first_col_bold=False,
)
add_note(doc, "Critical calendar action", "These five sessions replace the normal Saturday timetable on 31 October 2026. Suppress the recurring Saturday instances for that date before publishing the event classes.")

doc.add_heading("Booking Requirements", level=2)
for item in (
    "Online booking is mandatory even though the price is THB 0.",
    "Collect full name, phone number, email address, and preferred contact channel: WhatsApp, Facebook, or Instagram.",
    "Stop confirmed bookings automatically at 9 and enable a waitlist.",
    "Send a confirmation after booking and display the studio contact options, including WhatsApp, LINE, or phone.",
    "Save every booking as a customer contact in the back office.",
    "Keep the booking journey simple and mobile-friendly for tourists and local residents.",
):
    add_bullet(doc, item)

doc.add_heading("9 Customer and Booking Rules", level=1)
add_table(
    doc,
    ["Rule", "VMS configuration"],
    [
        ("Cancellation", "Free until 3 hours before start. Later cancellation or no-show deducts the class or session."),
        ("Refunds", "Packages are non-refundable."),
        ("Transfers", "Packages are non-transferable, except the defined class-sharing allowance."),
        ("Package validity", "Current group package validity begins on the purchase date."),
        ("New Reformer clients", "Ask clients to arrive 10 minutes early and tell the instructor they are new."),
        ("Grip socks", "Required for safety and hygiene."),
        ("Zero-booking class", "Internal rule: cancel at least 3 hours before start and notify the trainer by automated SMS."),
    ],
    widths=[4.1, 14.1],
    font_size=8.7,
    first_col_bold=True,
)

doc.add_page_break()
doc.add_heading("10 Internal Instructor Compensation", level=1)
add_note(doc, "Internal only", "Do not publish this section on customer-facing pages. Use it only if the VMS supports instructor payroll or session-cost rules.")
add_table(
    doc,
    ["Paid bookings", "Instructor pay"],
    [
        ("1-3 students", "THB 600"),
        ("4 students", "THB 700"),
        ("5 students", "THB 800"),
        ("6 students", "THB 900"),
        ("7 students", "THB 1,000"),
        ("8 students", "THB 1,100"),
        ("9 or more students", "THB 1,300 maximum"),
    ],
    widths=[8.8, 9.4],
    font_size=8.8,
    first_col_bold=True,
)
add_bullet(doc, "Late cancellations and no-shows count for instructor pay when the studio collects the booking fee.")
add_bullet(doc, "Private-class revenue split: 40 percent instructor and 60 percent studio, based on completed paid sessions.")
add_bullet(doc, "Planned trainer absence: 24 hours' notice and an approved substitute arranged by the trainer.")
add_bullet(doc, "Customer-refund complaints may trigger review, adjustment, or withholding of instructor compensation.")

doc.add_heading("11 Source Conflicts and Final VMS Actions", level=1)
add_table(
    doc,
    ["Field", "Latest client content", "Older material", "VMS action"],
    [
        ("Group validity start", "Purchase date", "First class attended", "Use purchase date"),
        ("10-class sharing", "3 classes", "2 classes", "Use 3"),
        ("30-class sharing", "15 classes", "10 classes", "Use 15"),
        ("Group massage gift", "Thai or foot massage", "Package-specific Thai or aroma", "Use Thai or foot"),
        ("Private 5-pack validity", "30 days", "60 days", "Use 30 days"),
        ("Private 10-pack", "Not included", "THB 20,900 for 90 days", "Hold for confirmation"),
        ("Two-person private promo", "Not included", "THB 2,500 promotional", "Hold for confirmation"),
        ("Private duration", "Not restated", "1 hour", "Confirm before publishing"),
    ],
    widths=[3.5, 4.7, 5.1, 4.9],
    font_size=8.0,
    first_col_bold=True,
)

doc.add_page_break()
doc.add_heading("12 Missing Information", level=1)
for item in (
    "Bophut private-session slot structure and confirmed class duration.",
    "Whether the legacy 10-session private package remains available.",
    "Whether the older two-person group-private promotion remains available.",
    "Stock limit for the complimentary massage gift.",
    "Trainer bios, qualifications, languages, contact details, and profile copy.",
    "Any Mat Yoga service, timetable, capacity, trainer allocation, and pricing. None was supplied in the current source set.",
    "Any memberships separate from class packages. None was supplied in the current source set.",
):
    add_bullet(doc, item)

doc.add_heading("13 Recommended VMS Build Order", level=1)
steps = [
    "Create the two studio areas and apply their capacities.",
    "Create the six regular group class types and the private service types.",
    "Create Kru Yui, Kru Nex, and Kru Leaw as trainer records.",
    "Publish the 32 recurring weekly group sessions and assign the correct trainer.",
    "Create the six group purchase options and current private purchase options.",
    "Configure cancellation, validity, sharing, and complimentary-massage rules.",
    "Create the five Halloween event classes as dated sessions and suppress the regular Saturday instances on 31 October.",
    "Enable capacity stops, waitlists, confirmations, contact collection, and customer-record creation.",
    "Add instructor compensation rules if payroll automation is available.",
    "Resolve the items in Missing Information before publishing the affected products.",
]
for i, text in enumerate(steps, 1):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    style_run(r, 9.6)

doc.add_heading("Sources Reviewed", level=1)
for source in (
    "Client message: Halloween soft-opening free classes and booking requirements",
    "Client message: Nathon private-service description and fixed service times",
    "Client message: current private-class pricing",
    "Client message marked final: group class packages",
    "Reformer Pilates Timetable versions 2 and 3",
    "Soft Opening Promo Menu",
    "Pilates Princess Soft Opening package menu",
    "Instructor Compensation Policy Bophut",
):
    add_bullet(doc, source)

# Document metadata
doc.core_properties.title = "MÀLEE Pilates VMS Master Data"
doc.core_properties.subject = "Consolidated services, timetable, trainers, pricing and launch promotions"
doc.core_properties.author = "Exsportia"
doc.core_properties.keywords = "MÀLEE Pilates, VMS, timetable, pricing, trainers, launch"

doc.save(OUT)
print(OUT)
